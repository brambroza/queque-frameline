#!/usr/bin/env python3
"""สร้างไฟล์สัญญา GoAlong–Fameline (.docx) จากไฟล์ร่าง .md ในโฟลเดอร์เดียวกัน

วิธีใช้:  python3 docs/contract/build_contract.py
- เนื้อหาแก้ที่ไฟล์ .md เท่านั้น แล้วรันสคริปต์นี้ใหม่
- ใช้ helper และสไตล์ชุดเดียวกับ docs/proposal/build_doc.py และ build_spec.py (ฟอนต์ไทย, ตาราง, ไฮไลต์ร่าง)
- ข้อความในวงเล็บเหลี่ยม เช่น [●] ถูกไฮไลต์สีเหลือง = ต้องกรอกหรือเคาะก่อนส่งลูกค้า
  (ตั้ง HIGHLIGHT_DRAFT = False ใน build_doc.py เพื่อเอาไฮไลต์ออก)
"""
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "proposal"))
sys.dont_write_bytecode = True  # import build_doc โดยไม่ทิ้ง .pyc ไว้ใน repo

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from build_doc import (  # noqa: E402
    BASE_PT, DARK, FONT, GREEN, GREY,
    _set_fonts, add_page_number, apply_widths, h2, note, page_break, para, run, shade,
)

SOURCE = HERE / "สัญญาจ้างพัฒนาซอฟต์แวร์-GoAlong-Fameline-ร่าง-v0.1.md"
OUTPUT = SOURCE.with_suffix(".docx")

PAGE_WIDTH_CM = 16.5
KEEP_TOGETHER_ROWS = 10
SIGN_SPACE_PT = 26
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# **ตัวหนา** หรือ [ข้อความที่ต้องกรอก]
INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\])")
SUB_CLAUSE = re.compile(r"^\([ก-ฮ0-9]+\)\s")
TABLE_DIVIDER = re.compile(r"^\|[\s:|-]+\|$")


def setup_page(doc):
    """ตั้งค่าหน้ากระดาษ A4, ฟอนต์ไทยของ style Normal และ footer เลขหน้า"""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BASE_PT)
    _set_fonts(normal.element.get_or_add_rPr(), size_pt=BASE_PT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "th-TH")
    normal.element.get_or_add_rPr().append(lang)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(footer, "GoAlong  |  สัญญาจ้างพัฒนาซอฟต์แวร์และให้บริการ — Fameline Dock Queue  |  หน้า ", size=8.5, color=GREY)
    add_page_number(footer)


def inline(par, text, *, size=None, bold=False, color=None):
    """เขียนข้อความหนึ่งบรรทัดลง paragraph โดยแปลง **ตัวหนา** และไฮไลต์ [ช่องที่ต้องกรอก]"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            inline(par, part[2:-2], size=size, bold=True, color=color)
        elif part.startswith("["):
            run(par, part, size=size, bold=bold, draft=True)
        else:
            run(par, part, size=size, bold=bold, color=color)


def heading(doc, text):
    """หัวข้อระดับข้อสัญญา — สไตล์เดียวกับ h1 ของ build_doc (เขียว + เส้นใต้) แต่ใช้เลขข้อจากต้นฉบับ"""
    p = para(doc, before=14, after=8)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, size=13, color=GREEN)
    ppr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", "0B7A4B")):
        bottom.set(qn(k), v)
    border.append(bottom)
    ppr.append(border)


def title(doc, text, *, first):
    """ชื่อเอกสาร / ชื่อเอกสารแนบท้าย — เอกสารแนบขึ้นหน้าใหม่เสมอ"""
    if not first:
        page_break(doc)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    run(p, text, bold=True, size=17 if first else 15, color=GREEN)


def column_widths(rows):
    """แบ่งความกว้างคอลัมน์ตามความยาวข้อความที่ยาวที่สุดของแต่ละคอลัมน์ (มีขั้นต่ำและเพดาน)"""
    weights = [min(max(max(len(r[i]) for r in rows), 8), 60) for i in range(len(rows[0]))]
    total = sum(weights)
    return [PAGE_WIDTH_CM * w / total for w in weights]


def table(doc, rows):
    """ตารางจาก markdown — หัวเขียว แถวสลับสี เหมือน table() ของ build_doc แต่รองรับตัวหนาและช่องที่ต้องกรอก"""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = t.rows[ri].cells[ci]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            # ตารางสั้น (เช่น ช่องลงนาม) ต้องไม่ถูกตัดข้ามหน้า
            p.paragraph_format.keep_with_next = len(rows) <= KEEP_TOGETHER_ROWS and ri < len(rows) - 1
            if ri == 0:
                shade(cell, "0B7A4B")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run(p, text.replace("**", ""), bold=True, size=9.5, color=WHITE)
                continue
            if text.startswith("ลงชื่อ"):
                p.paragraph_format.space_before = Pt(SIGN_SPACE_PT)  # เว้นที่ให้เซ็น
            if ri % 2 == 0:
                shade(cell, "F3F7F5")
            inline(p, text, size=9.5)
    apply_widths(t, column_widths(rows))
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)
    para(doc, after=4)


def split_row(line):
    """แยกบรรทัดตาราง markdown เป็นรายการข้อความของแต่ละช่อง"""
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def body_line(doc, text):
    """ย่อหน้าเนื้อความ — ข้อย่อย (ก) (ข) เยื้องเข้า, บรรทัดที่ขึ้นด้วย [ = ทางเลือกที่รอเคาะ"""
    p = para(doc, after=5)
    if SUB_CLAUSE.match(text):
        p.paragraph_format.left_indent = Cm(0.9)
    inline(p, text)


def bullet(doc, text):
    """รายการแบบ bullet สไตล์เดียวกับ bullets() ของ build_doc"""
    p = para(doc, after=3)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    run(p, "•  ", color=GREEN, bold=True)
    inline(p, text)


def build():
    """อ่านไฟล์ .md ทีละบรรทัด แปลงเป็น .docx แล้วบันทึก"""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_page(doc)

    first_title = True
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line or line == "---":
            continue
        if line.startswith("# "):
            title(doc, line[2:], first=first_title)
            first_title = False
        elif line.startswith("## ("):
            para(doc, line[3:], bold=True, size=12.5, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
        elif line.startswith("## "):
            heading(doc, line[3:])
        elif line.startswith("### "):
            h2(doc, line[4:])
        elif line.startswith("|"):
            rows = [split_row(line)]
            while i < len(lines) and lines[i].startswith("|"):
                if not TABLE_DIVIDER.match(lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            table(doc, rows)
        elif line.startswith("> "):
            note(doc, line[2:])
        elif line.startswith("- "):
            bullet(doc, line[2:])
        else:
            body_line(doc, line)

    doc.save(OUTPUT)
    print(f"saved {OUTPUT}")


if __name__ == "__main__":
    build()
